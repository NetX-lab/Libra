from pathlib import Path
import subprocess
import tempfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt


ROOT = Path("/Users/kevin/Documents/System Implementations/RL_Framework_npu")
PDF = ROOT / "Libra_NPU_Support_仿真工具报告.pdf"
OUTPUT = ROOT / "Libra_NPU_Support_仿真工具报告.docx"


def build():
    with tempfile.TemporaryDirectory(prefix="libra-npu-docx-pages.") as temp:
        prefix = str(Path(temp) / "page")
        subprocess.run(
            ["pdftoppm", "-png", "-r", "150", str(PDF), prefix],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        pages = sorted(Path(temp).glob("page-*.png"))
        document = Document()
        section = document.sections[0]
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0)
        section.bottom_margin = Inches(0)
        section.left_margin = Inches(0)
        section.right_margin = Inches(0)
        section.header_distance = Inches(0)
        section.footer_distance = Inches(0)
        document.core_properties.title = "Libra NPU_Support 分支仿真工具报告"
        document.core_properties.subject = "Agentic RL 后训练资源规划、成本建模与动态重规划"
        document.core_properties.author = "OpenAI Codex"
        for index, page in enumerate(pages):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1
            paragraph.paragraph_format.page_break_before = index > 0
            run = paragraph.add_run()
            run.add_picture(str(page), width=Inches(8.27), height=Inches(11.68))
        document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
