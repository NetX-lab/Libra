from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
MD_PATH = ROOT / "Libra_NPU_Support_仿真工具报告.md"
OUT_PATH = ROOT / "Libra_NPU_Support_仿真工具报告.docx"

NAVY = "1F4E78"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "D9E1F2"
TEXT = "222222"


def set_run_font(run, name="Aptos", size=10.5, bold=None, color=TEXT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    rfonts = run._element.get_or_add_rPr().rFonts
    rfonts.set(qn("w:eastAsia"), "Heiti SC")
    rfonts.set(qn("w:cs"), "Heiti SC")
    rfonts.set(qn("w:hint"), "eastAsia")
    lang = run._element.get_or_add_rPr().find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        run._element.get_or_add_rPr().append(lang)
    lang.set(qn("w:eastAsia"), "zh-CN")
    lang.set(qn("w:val"), "en-US")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color="666666")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Aptos")
    r_fonts.set(qn("w:hAnsi"), "Aptos")
    r_fonts.set(qn("w:eastAsia"), "Heiti SC")
    r_fonts.set(qn("w:cs"), "Heiti SC")
    r_pr.append(r_fonts)
    r.append(r_pr)
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_run_font(run, size=9, color="666666")


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Aptos")
    rfonts.set(qn("w:hAnsi"), "Aptos")
    rfonts.set(qn("w:eastAsia"), "Heiti SC")
    rfonts.set(qn("w:cs"), "Heiti SC")
    rpr.append(rfonts)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text, size=10.5):
    pattern = re.compile(r"(\[[^\]]+\]\([^\)]+\)|\*\*[^*]+\*\*|`[^`]+`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, size=size)
        token = match.group(0)
        if token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^\)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Aptos Mono", size=size - 0.5, color="444444")
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=size)


def style_document(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.85)
    sec.right_margin = Inches(0.85)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti SC")
    normal._element.rPr.rFonts.set(qn("w:cs"), "Heiti SC")
    normal._element.rPr.append(OxmlElement("w:lang"))
    normal._element.rPr[-1].set(qn("w:eastAsia"), "zh-CN")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    for name, size, space_before, space_after in [
        ("Title", 25, 0, 14),
        ("Heading 1", 16, 16, 7),
        ("Heading 2", 13, 11, 5),
        ("Heading 3", 11.5, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti SC")
        style._element.rPr.rFonts.set(qn("w:cs"), "Heiti SC")
        style._element.rPr.append(OxmlElement("w:lang"))
        style._element.rPr[-1].set(qn("w:eastAsia"), "zh-CN")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("000000")
        style.paragraph_format.space_before = Pt(space_before)
        style.paragraph_format.space_after = Pt(space_after)
        style.paragraph_format.keep_with_next = True

    # Remove the built-in Title style's theme-colored bottom border.
    title_style = styles["Title"]
    ppr = title_style._element.get_or_add_pPr()
    p_bdr = ppr.find(qn("w:pBdr"))
    if p_bdr is not None:
        ppr.remove(p_bdr)

    footer = sec.footer
    add_page_number(footer.paragraphs[0])


def add_para(doc, text, style=None, align=None, size=10.5):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.widow_control = True
    if style is None:
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.18
    add_inline(p, text, size=size)
    return p


def add_table(doc, lines):
    rows = []
    for line in lines:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append(cells)
    if len(rows) < 2:
        return
    headers = rows[0]
    body = [r for r in rows[2:] if any(r)]
    cols = len(headers)
    table = doc.add_table(rows=1, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    header_row = table.rows[0]
    repeat_table_header(header_row)
    for idx, value in enumerate(headers):
        cell = header_row.cells[idx]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(value)
        set_run_font(run, size=9, bold=True, color="FFFFFF")
    for row_idx, values in enumerate(body):
        while len(values) < cols:
            values.append("")
        row = table.add_row()
        for idx in range(cols):
            cell = row.cells[idx]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_idx % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            add_inline(p, values[idx], size=8.8)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_code(doc, code):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="C9D3E0", size="4")
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F6F8")
    set_cell_margins(cell, top=130, start=170, bottom=130, end=170)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    for idx, line in enumerate(code.splitlines()):
        run = p.add_run(line)
        set_run_font(run, name="Aptos Mono", size=8.6, color="333333")
        if idx < len(code.splitlines()) - 1:
            run.add_break()
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def parse_markdown(doc, text):
    lines = text.splitlines()
    i = 0
    h2_titles = []
    while i < len(lines):
        line = lines[i]
        if line.startswith("# "):
            i += 1
            continue
        if line.startswith("## "):
            title = line[3:].strip()
            h2_titles.append(title)
            add_para(doc, title, style="Heading 1")
            i += 1
            continue
        if line.startswith("### "):
            add_para(doc, line[4:].strip(), style="Heading 2")
            i += 1
            continue
        if line.startswith("#### "):
            add_para(doc, line[5:].strip(), style="Heading 3")
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        if line.startswith("```"):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code_lines.append(lines[i])
                i += 1
            if i < len(lines):
                i += 1
            add_code(doc, "\n".join(code_lines))
            continue
        if line.startswith("|") and i + 1 < len(lines) and lines[i + 1].startswith("|"):
            table_lines = [line, lines[i + 1]]
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_table(doc, table_lines)
            continue
        if re.match(r"^\s*[-*] ", line):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, re.sub(r"^\s*[-*] ", "", line), size=10.2)
            i += 1
            continue
        if re.match(r"^\s*\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, re.sub(r"^\s*\d+\. ", "", line), size=10.2)
            i += 1
            continue
        para_lines = [line.strip()]
        i += 1
        while i < len(lines):
            nxt = lines[i]
            if (not nxt.strip() or nxt.startswith(("# ", "## ", "### ", "#### ", "```", "|")) or re.match(r"^\s*(?:[-*] |\d+\. )", nxt)):
                break
            para_lines.append(nxt.strip())
            i += 1
        add_para(doc, " ".join(para_lines))
    return h2_titles


def main():
    doc = Document()
    style_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Libra NPU Support 分支仿真工具报告")
    set_run_font(run, name="Aptos Display", size=25, bold=True, color="000000")
    title_ppr = title._p.get_or_add_pPr()
    title_pbdr = title_ppr.find(qn("w:pBdr"))
    if title_pbdr is not None:
        title_ppr.remove(title_pbdr)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("面向 Ascend NPU 异步 Agentic RL 资源规划")
    set_run_font(run, size=13, color="4F6475")
    doc.add_paragraph()
    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(meta, color="D9D9D9", size="5")
    values = [
        ("分析对象", "NetX-lab/Libra NPU_Support 分支"),
        ("代码基线", "365f5e42d23b8641ac32e1adaa6feac6b5aa812b"),
        ("报告日期", "2026-09-09"),
        ("验证范围", "源码审查、CPU 仿真流程、CPU 友好测试"),
    ]
    for ridx, (k, v) in enumerate(values):
        for cidx, value in enumerate((k, v)):
            cell = meta.cell(ridx, cidx)
            set_cell_margins(cell, top=120, start=150, bottom=120, end=150)
            if cidx == 0:
                set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(value)
            set_run_font(run, size=9.5, bold=(cidx == 0))
    doc.add_paragraph()
    lead = doc.add_paragraph()
    lead.alignment = WD_ALIGN_PARAGRAPH.LEFT
    lead.paragraph_format.space_before = Pt(18)
    lead.paragraph_format.space_after = Pt(8)
    run = lead.add_run("主要结论  ")
    set_run_font(run, size=11, bold=True, color=NAVY)
    add_inline(lead, "该分支已形成面向 Ascend NPU 的资源规划仿真路径，但当前验证结果主要证明软件流程和规划逻辑，真实 NPU 性能仍需通过 CANN、HCCL 和 vLLM-Ascend 多节点实测校准。", size=11)
    doc.add_page_break()

    # Manual contents page for stable rendering without Word field updates.
    add_para(doc, "目录", style="Heading 1")
    contents = [
        "报告摘要",
        "1 项目定位与分析范围",
        "2 仿真工具总体架构",
        "3 仿真对象与输入输出",
        "4 内置解析成本模型",
        "5 Sailor 和 Vidur 适配器",
        "6 GRP 搜索和动态重规划",
        "7 NPU 运行链路与仿真关系",
        "8 安装和运行方法",
        "9 验证结果",
        "10 输出文件和可观测性",
        "11 适用性评估",
        "12 改进建议",
        "13 总结",
        "14 参考资料",
    ]
    for entry in contents:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(entry)
        set_run_font(run, size=10.5)
    doc.add_page_break()

    parse_markdown(doc, MD_PATH.read_text(encoding="utf-8"))
    doc.core_properties.title = "Libra NPU Support 分支仿真工具报告"
    doc.core_properties.subject = "Ascend NPU 异步 Agentic RL 资源规划仿真工具"
    doc.core_properties.author = "Codex"
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    main()
